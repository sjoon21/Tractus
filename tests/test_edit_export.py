"""Phase 5 CLI: 세그먼트 텍스트 편집, 용어집 추가, 내보내기. 회의 폴더는 save_transcript로 직접 만든다 (모델·FFmpeg 불필요)."""

import csv
import io
import json
import zipfile
from pathlib import Path

import pytest
from typer.testing import CliRunner

from tractus import cli
from tractus.store import load_glossary, render_vtt, save_transcript

MID = "01K6X3TESTMEETING000000000"


def _transcript() -> dict:
    return {
        "schema_version": 2, "meeting_id": MID, "title": "주간회의", "created_at": "2026-10-02T14:00:00+09:00",
        "audio": {"file": "audio.ogg", "duration_sec": 3702.0, "source": "주간회의.m4a", "sha256": "x"},
        "asr": {"model": "m", "track": "lite", "runtime": "mlx-audio", "language": "auto", "hotwords": []},
        "speakers": {"S1": {"name": "김민수"}, "S2": {"name": None}},
        "diarization": None,
        "segments": [
            {"id": "s0001", "start": 0.0, "end": 2.5, "text": "안녕하세요", "lang": "ko", "speaker": "S1", "corrected_text": "안녕하십니까", "words": None},
            {"id": "s0002", "start": 3.0, "end": 65.123, "text": "예산, 보겠습니다", "lang": "ko", "speaker": "S2", "corrected_text": None, "words": None},
            {"id": "s0003", "start": 3700.5, "end": 3702.0, "text": "마무리", "lang": "ko", "speaker": None, "corrected_text": None, "words": None},
        ],
    }


SUMMARY = {
    "schema_version": 1, "meeting_id": MID, "provider": "upstage", "model": "m", "generated_at": "g",
    "title": "주간회의 요약", "summary": "예산을 검토했다.", "topics": [{"title": "예산", "segment_ids": ["s0002"]}],
    "decisions": [{"text": "예산 동결", "segment_ids": ["s0002"]}],
    "action_items": [{"task": "견적 받기", "owner": "김민수", "due": "금요일", "segment_ids": ["s0001"]}],
}


@pytest.fixture
def meeting(tmp_path, monkeypatch):
    home = tmp_path / "home"
    d = home / "meetings" / f"2026-10-02_주간회의_{MID}"
    d.mkdir(parents=True)
    save_transcript(d, _transcript())
    (d / "audio.ogg").write_bytes(b"OggS fake audio")
    (d / "summary.md").write_text("# 주간회의 요약\n\n예산을 검토했다.\n", encoding="utf-8")
    (d / "summary.json").write_text(json.dumps(SUMMARY, ensure_ascii=False), encoding="utf-8")
    (d / "memos.json").write_text('[{"t": 1.0, "kind": "memo", "text": "메모"}]', encoding="utf-8")
    (d / "tracks").mkdir()
    (d / "tracks" / "mic.wav").write_bytes(b"RIFF")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("TRACTUS_HOME", str(home))
    return d


def invoke(*args, input=None):
    return CliRunner().invoke(cli.app, [*map(str, args)], input=input)


def read_t(d: Path) -> dict:
    return json.loads((d / "transcript.json").read_text(encoding="utf-8"))


# ── segments edit ───────────────────────────────────────────────────


def test_segments_edit_writes_renders_and_syncs_without_embedding(meeting, monkeypatch):
    synced = []
    monkeypatch.setattr(cli, "_sync_index_quietly", lambda embed=True: synced.append(embed))
    edits = {"edits": [{"id": "s0002", "text": " 예산을 보겠습니다 "}, {"id": "s0001", "text": "안녕하세요"}]}
    r = invoke("segments", MID[-6:], "edit", input=json.dumps(edits, ensure_ascii=False))
    assert r.exit_code == 0, r.output
    assert json.loads(r.stdout) == {"changed": 2}
    segs = read_t(meeting)["segments"]
    assert [s["corrected_text"] for s in segs] == [None, "예산을 보겠습니다", None]  # s0001은 원문과 같아 교정 해제
    assert [s["text"] for s in segs] == ["안녕하세요", "예산, 보겠습니다", "마무리"] and segs[1]["words"] is None
    assert "**참석자 2**: 예산을 보겠습니다" in (meeting / "transcript.md").read_text(encoding="utf-8")
    assert "<v 참석자 2>예산을 보겠습니다" in (meeting / "transcript.vtt").read_text(encoding="utf-8")
    assert synced == [False] and not list(meeting.glob("*.tmp"))


@pytest.mark.parametrize("payload", [
    '{"edits": [{"id": "s0001", "text": "x"}, {"id": "s9999", "text": "y"}]}',
    '{"edits": [{"id": "s0001", "text": "   "}]}',
    '{"edits": "s0001"}',
    'not json',
])
def test_segments_edit_rejects_bad_input_and_writes_nothing(meeting, monkeypatch, payload):
    monkeypatch.setattr(cli, "_sync_index_quietly", lambda embed=True: pytest.fail("색인하면 안 됨"))
    before = {p.name: p.read_bytes() for p in meeting.glob("transcript.*")}
    r = invoke("segments", str(meeting), "edit", input=payload)
    assert r.exit_code == 1 and r.stdout == "" and r.stderr.strip()
    assert {p.name: p.read_bytes() for p in meeting.glob("transcript.*")} == before


def test_segments_edit_unknown_id_names_it(meeting):
    r = invoke("segments", MID, "edit", input='{"edits": [{"id": "s9999", "text": "y"}]}')
    assert r.exit_code == 1 and r.stderr.strip().splitlines()[-1] == "해당 대화를 찾지 못했습니다. 노트를 다시 열어 주세요."


# ── glossary add ────────────────────────────────────────────────────


def test_glossary_add_creates_file_and_dedupes(meeting):
    path = meeting.parents[1] / "glossary.md"
    r = invoke("glossary", "add", " 쿠버네티스 ", "WTO", "", "쿠버네티스")
    assert r.exit_code == 0, r.output
    assert json.loads(r.stdout) == {"added": 2} and load_glossary(path) == ["쿠버네티스", "WTO"]
    r = invoke("glossary", "add", "WTO", "Project Atlas")
    assert json.loads(r.stdout) == {"added": 1} and load_glossary(path)[-1] == "Project Atlas"
    r = invoke("glossary", "add", "--", "-x")  # 앱은 용어 앞에 --를 둔다
    assert json.loads(r.stdout) == {"added": 1} and load_glossary(path)[-1] == "-x"


# ── export ──────────────────────────────────────────────────────────


def export(fmt, out, *extra):
    r = invoke("export", MID, "--format", fmt, "--out", out, *extra)
    assert r.exit_code == 0, r.output
    assert json.loads(r.stdout) == {"path": str(Path(out).resolve())}
    return Path(out)


def test_export_txt_uses_display_names_and_corrected_text(meeting, tmp_path):
    lines = export("txt", tmp_path / "o.txt").read_text(encoding="utf-8").splitlines()
    assert lines[0] == "주간회의" and "2026-10-02 14:00" in lines[1] and "01:01:42" in lines[1]
    assert lines[-3:] == ["[00:00:00] 김민수: 안녕하십니까", "[00:00:03] 참석자 2: 예산, 보겠습니다", "[01:01:40] 마무리"]


def test_export_md_appends_summary(meeting, tmp_path):
    text = export("md", tmp_path / "o.md").read_text(encoding="utf-8")
    assert "[00:00:00] **김민수**: 안녕하십니까" in text and "(메모) 메모" in text
    assert text.index("**김민수**") < text.index("# 주간회의 요약")


def test_export_srt_timing_and_vtt(meeting, tmp_path):
    srt = export("srt", tmp_path / "o.srt").read_text(encoding="utf-8")
    assert srt.startswith(
        "1\n00:00:00,000 --> 00:00:02,500\n김민수: 안녕하십니까\n\n"
        "2\n00:00:03,000 --> 00:01:05,123\n참석자 2: 예산, 보겠습니다\n\n"
        "3\n01:01:40,500 --> 01:01:42,000\n마무리\n"
    )
    assert export("vtt", tmp_path / "o.vtt").read_text(encoding="utf-8") == render_vtt(read_t(meeting))


def test_export_csv_is_utf8_sig_with_korean_headers(meeting, tmp_path):
    raw = export("csv", tmp_path / "o.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # 한국어 Excel이 UTF-8로 연다
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0] == ["시작", "끝", "참석자", "내용"]
    assert rows[1:] == [["00:00:00", "00:00:02", "김민수", "안녕하십니까"], ["00:00:03", "00:01:05", "참석자 2", "예산, 보겠습니다"], ["01:01:40", "01:01:42", "", "마무리"]]


def test_export_docx_roundtrip(meeting, tmp_path):
    import docx
    from docx.oxml.ns import qn

    doc = docx.Document(str(export("docx", tmp_path / "o.docx")))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "주간회의"
    body = "\n".join(texts)
    assert "2026-10-02 14:00" in body and "01:01:42" in body and "참석자 2명 (김민수, 참석자 2)" in body and "대화 내용" in texts
    assert "예산을 검토했다." in body and "예산 동결" in body and "견적 받기 — 김민수, 금요일" in body
    line = next(p for p in doc.paragraphs if "안녕하십니까" in p.text)
    assert line.text == "00:00:00  김민수  안녕하십니까"
    assert [r.text for r in line.runs if r.bold] == ["김민수"]
    assert doc.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == "Apple SD Gothic Neo"


def test_export_docx_without_summary(meeting, tmp_path):
    import docx

    (meeting / "summary.json").unlink()
    body = "\n".join(p.text for p in docx.Document(str(export("docx", tmp_path / "o.docx"))).paragraphs)
    assert "결정 사항" not in body and "마무리" in body


def test_export_bundle_zips_meeting_files_only(meeting, tmp_path):
    with zipfile.ZipFile(export("bundle", tmp_path / "o.zip")) as z:
        names = sorted(z.namelist())
        assert z.read(f"{meeting.name}/audio.ogg") == b"OggS fake audio"
    assert names == sorted(f"{meeting.name}/{n}" for n in (
        "audio.ogg", "transcript.json", "transcript.md", "transcript.vtt", "summary.md", "summary.json", "memos.json"))


def test_export_refuses_to_overwrite_without_force(meeting, tmp_path):
    out = tmp_path / "o.txt"
    out.write_text("keep", encoding="utf-8")
    r = invoke("export", MID, "--format", "txt", "--out", out)
    assert r.exit_code == 1 and out.read_text(encoding="utf-8") == "keep" and "--force" in r.stderr
    assert "김민수" in export("txt", out, "--force").read_text(encoding="utf-8")
    assert invoke("export", MID, "--format", "pdf", "--out", tmp_path / "o.pdf").exit_code != 0


# ── 통합 점검·리뷰 반영 ─────────────────────────────────────────────


def test_long_error_stays_on_one_line_and_names_out(meeting, tmp_path):
    out = tmp_path / ("긴경로" * 20) / "없는폴더" / "x.txt"  # 앱은 stderr 마지막 줄만 보여 준다
    r = invoke("export", MID, "--format", "txt", "--out", out, "--force")
    assert r.exit_code == 1
    last = r.stderr.strip().splitlines()[-1]
    assert last.startswith(str(out)) and last.endswith(")") and "에 저장할 수 없습니다" in last and ".tmp" not in last


def test_control_chars_are_dropped_from_edits_and_names(meeting, monkeypatch):
    monkeypatch.setattr(cli, "_sync_index_quietly", lambda embed=True: None)
    r = invoke("segments", MID, "edit", input=json.dumps({"edits": [{"id": "s0002", "text": "abc\u0001def\u0085g"}]}))
    assert r.exit_code == 0, r.output
    assert invoke("speakers", MID, "rename", "S2", "이\x02영희").exit_code == 0
    t = read_t(meeting)
    assert t["segments"][1]["corrected_text"] == "abcdef g" and t["speakers"]["S2"]["name"] == "이영희"


def test_docx_export_error_is_one_line_not_traceback(meeting, tmp_path):
    t = read_t(meeting)
    save_transcript(meeting, {**t, "title": "제목\x01"})
    r = invoke("export", MID, "--format", "docx", "--out", tmp_path / "o.docx")
    lines = r.stderr.strip().splitlines()
    assert r.exit_code == 1 and "Traceback" not in r.output and lines[-1] == "제목이나 대화 내용에 쓸 수 없는 글자가 있습니다."
    assert any(line.startswith("All strings must be XML compatible") for line in lines[:-1])  # 원인은 앞 줄 (CLI용)


def test_csv_neutralizes_formula_cells(meeting, tmp_path):
    t = read_t(meeting)
    t["segments"][1]["corrected_text"] = "=1+2"
    t["speakers"]["S1"]["name"] = "@이름"
    save_transcript(meeting, t)
    rows = list(csv.reader(io.StringIO(export("csv", tmp_path / "o.csv").read_bytes().decode("utf-8-sig"))))
    assert rows[1][2] == "'@이름" and rows[2][3] == "'=1+2" and rows[3][3] == "마무리"


def test_docx_headings_use_east_asian_font(meeting, tmp_path):
    import docx
    from docx.oxml.ns import qn

    doc = docx.Document(str(export("docx", tmp_path / "o.docx")))
    for name in ("Title", "Heading 1"):
        rf = doc.styles[name].element.rPr.rFonts
        assert rf.get(qn("w:eastAsia")) == "Apple SD Gothic Neo" and rf.get(qn("w:eastAsiaTheme")) is None


def test_md_export_drops_summary_front_matter(meeting, tmp_path):
    (meeting / "summary.md").write_text("---\nmeeting_id: x\nprovider: upstage\n---\n\n# 주간회의 요약\n", encoding="utf-8")
    text = export("md", tmp_path / "o.md").read_text(encoding="utf-8")
    assert "provider: upstage" not in text and "# 주간회의 요약" in text


def test_srt_and_txt_keep_one_line_per_cue(meeting, tmp_path):
    t = read_t(meeting)
    t["segments"][2]["text"] = "a --> b\n\n줄  바꿈"
    save_transcript(meeting, t)
    assert export("srt", tmp_path / "o.srt").read_text(encoding="utf-8").endswith("01:01:40,500 --> 01:01:42,000\na → b 줄 바꿈\n")
    assert export("txt", tmp_path / "o.txt").read_text(encoding="utf-8").endswith("[01:01:40] a --> b 줄 바꿈\n")


# ── 목소리 제안은 내보내지 않는다 (Phase 7.1) ───────────────────────


def test_voice_suggestion_never_leaks_into_exports(meeting, tmp_path):
    import docx

    t = read_t(meeting)
    t["speakers"]["S2"] = {"name": None, "suggested": {"person_id": "P1", "name": "홍길동", "score": 0.82}}
    save_transcript(meeting, t)
    for fmt in ("txt", "md", "srt", "vtt"):
        assert "홍길동" not in export(fmt, tmp_path / f"o.{fmt}").read_text(encoding="utf-8"), fmt
    assert "홍길동" not in export("csv", tmp_path / "o.csv").read_bytes().decode("utf-8-sig")
    assert "홍길동" not in "\n".join(p.text for p in docx.Document(str(export("docx", tmp_path / "o.docx"))).paragraphs)
    with zipfile.ZipFile(export("bundle", tmp_path / "o.zip")) as z:
        raw = {n: z.read(n).decode("utf-8", "replace") for n in z.namelist()}
    assert not any("홍길동" in v or "suggested" in v for v in raw.values())
    assert json.loads(raw[f"{meeting.name}/transcript.json"])["speakers"] == {"S1": {"name": "김민수"}, "S2": {"name": None}}
    assert "suggested" in (meeting / "transcript.json").read_text(encoding="utf-8")  # 원본은 그대로


def test_summary_prompt_and_owners_never_use_suggested_name():
    from tractus.summarize import ActionItem, MeetingSummary, build_prompt, clean_owners

    t = _transcript()
    t["speakers"]["S2"] = {"name": None, "suggested": {"person_id": "P1", "name": "홍길동", "score": 0.82}}
    assert "홍길동" not in build_prompt(t) and " 참석자 2] " in build_prompt(t)
    s = MeetingSummary(title="t", summary="s", topics=[], decisions=[], action_items=[
        ActionItem(task="a", owner="홍길동", due=None, segment_ids=["s0002"]), ActionItem(task="b", owner="김민수", due=None, segment_ids=["s0001"])])
    assert [a.owner for a in clean_owners(s, t).action_items] == [None, "김민수"]


# ── export: 넣을 내용 고르기 (--summary 여러 번 + --parts) ─────────────────────

MEMO_BLOCK = {"body": "첫 메모\n\n둘째", "lines": [{"id": "m0001", "line": 1, "text": "첫 메모"}, {"id": "m0002", "line": 3, "text": "둘째"}], "seq": 2}


def _docx_paras(path: Path) -> list[tuple[str, str]]:
    import docx

    return [(p.style.name, p.text) for p in docx.Document(str(path)).paragraphs]


def _write_v2(d: Path, key: str, data: dict) -> None:
    from tractus.summaries import render_md

    data = {"schema_version": 2, "meeting_id": MID, "template": key, "provider": "p", "model": "m", "generated_at": "g", "overview": "", **data}
    (d / "summaries").mkdir(exist_ok=True)
    (d / "summaries" / f"{key}.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    (d / "summaries" / f"{key}.md").write_text(render_md(data, read_t(d)), encoding="utf-8")


def _line(text: str, children: list | None = None) -> dict:
    return {"text": text, "segment_ids": ["s0002"], **({"children": children} if children is not None else {})}


@pytest.fixture
def full(meeting):
    """기본 요약 + 핵심 요약 + 회의록 + 노트 완성 + 메모 탭 메모 + 녹음 중 메모·북마크가 모두 있는 노트."""
    from tractus.minutes import to_sections

    save_transcript(meeting, {**_transcript(), "memo": MEMO_BLOCK})
    (meeting / "summary.md").write_text("---\nmeeting_id: x\n---\n\n# 주간회의 요약\n\n예산을 검토했다.\n", encoding="utf-8")
    (meeting / "memos.json").write_text('[{"t": 1.0, "kind": "memo", "text": "메모"}, {"t": 70.0, "kind": "bookmark"}]', encoding="utf-8")
    _write_v2(meeting, "brief", {"template_name": "핵심 요약", "title": "핵심 요약 제목", "overview": "짧은 개요",
                                 "sections": [{"heading": "핵심 포인트", "style": "bullets", "items": [{**_line("포인트 하나"), "owner": None, "due": None}]}]})
    m = {"agenda": "예산 검토", "next_meeting": [_line("다음 주 월요일")],
         "topics": [{"title": "예산", "open": False, "discussed": [_line("예산 논의")], "agreed": [_line("동결")], "approach": [_line("분기별로 본다")],
                     "points": [{"label": "현황", "details": [_line("세부 줄", [_line("하위 줄")])]}]},
                    {"title": "채용", "open": True, "discussed": [], "agreed": [], "approach": [], "points": [{"label": "인원", "details": [_line("두 명", [])]}]}],
         "actions": [{**_line("견적 받기"), "team": "재무팀", "owner": "김민수", "due": "금요일", "note": None}]}
    _write_v2(meeting, "minutes", {"kind": "minutes", "template_name": "회의록", "title": "주간회의", "overview": m["agenda"], "minutes": m,
                                   "attendees": ["김민수", "이영희"], "edited_at": None, "sections": to_sections(m)})
    _write_v2(meeting, "enhanced", {"kind": "enhanced", "template_name": "노트 완성", "title": "완성 제목",
                                    "sections": [{"heading": "첫 메모", "style": "bullets", "part": "memo", "source": "m0001", "items": []}]})
    return meeting


def test_export_without_new_flags_is_byte_identical(full, tmp_path):
    md = (
        "---\nmeeting_id: 01K6X3TESTMEETING000000000\ntitle: 주간회의\ncreated_at: 2026-10-02T14:00:00+09:00\nduration: 01:01:42\nasr_model: m\n---\n\n# 주간회의\n\n"
        "[00:00:00] **김민수**: 안녕하십니까\n[00:00:01] (메모) 메모\n[00:00:03] **참석자 2**: 예산, 보겠습니다\n[00:01:10] (북마크)\n[01:01:40] 마무리\n\n"
        "## 메모\n\n첫 메모\n\n둘째\n\n---\n\n# 주간회의 요약\n\n예산을 검토했다.\n"
    )
    assert export("md", tmp_path / "a.md").read_text(encoding="utf-8") == md
    assert export("md", tmp_path / "b.md", "--summary", "meeting").read_text(encoding="utf-8") == md
    paras = [
        ("Title", "주간회의"), ("Normal", "일시 2026-10-02 14:00 · 길이 01:01:42"), ("Normal", "참석자 2명 (김민수, 참석자 2)"),
        ("Heading 1", "요약"), ("Normal", "예산을 검토했다."), ("Heading 1", "결정 사항"), ("List Bullet", "예산 동결"), ("Heading 1", "할 일"),
        ("List Bullet", "견적 받기 — 김민수, 금요일"), ("Heading 1", "대화 내용"), ("Normal", "00:00:00  김민수  안녕하십니까"),
        ("Normal", "00:00:03  참석자 2  예산, 보겠습니다"), ("Normal", "01:01:40  마무리"), ("Heading 1", "메모"), ("Normal", "첫 메모"), ("Normal", "둘째"),
    ]
    assert _docx_paras(export("docx", tmp_path / "a.docx")) == paras


def test_export_selected_summaries_and_parts_in_order(full, tmp_path):
    sel = ["--summary", "minutes", "--summary", "enhanced", "--summary", "meeting", "--summary", "brief", "--parts", "marks,memo,transcript"]
    md = export("md", tmp_path / "o.md", *sel).read_text(encoding="utf-8")
    order = ["# 주간회의\n", "# [26.10.02] 주간회의", "# 완성 제목", "# 주간회의 요약", "# 핵심 요약 제목", "## 대화 내용", "## 메모\n", "## 녹음 중 메모"]
    assert [md.index(x) for x in order] == sorted(md.index(x) for x in order)
    assert "meeting_id: x" not in md and "provider: p" not in md and md.count("\n---\n") == 8  # 요약의 frontmatter는 뺀다 (머리 frontmatter 끝 1 + 부분 사이 7)
    assert md.endswith("## 녹음 중 메모\n\n[00:00:01] (메모) 메모\n[00:01:10] (북마크)\n")
    assert "[00:00:03] **참석자 2**: 예산, 보겠습니다\n[01:01:40] 마무리" in md  # 녹음 중 메모는 대화 사이에 끼우지 않는다

    paras = _docx_paras(export("docx", tmp_path / "o.docx", *sel))
    texts = [x for _, x in paras]
    heads = ["요약 · 회의록", "요약 · 노트 완성", "요약", "요약 · 핵심 요약", "대화 내용", "메모", "녹음 중 메모"]
    at = [next(i for i, p in enumerate(paras) if p == ("Heading 1", h)) for h in heads]
    assert at == sorted(at) and texts[-2:] == ["00:00:01  (메모) 메모", "00:01:10  (북마크)"]


def test_minutes_docx_keeps_details_and_points_only_topics(full, tmp_path):
    paras = _docx_paras(export("docx", tmp_path / "o.docx", "--summary", "minutes", "--parts", ""))
    assert ("Heading 1", "1) 예산") in paras and ("Heading 1", "2) [추가 논의 필요] 채용") in paras  # 이름표만 있는 안건도 빠지지 않는다
    assert ("List Bullet", "현황") in paras and ("List Bullet 2", "세부 줄") in paras and ("List Bullet 3", "하위 줄") in paras and ("List Bullet 2", "두 명") in paras
    texts = [x for _, x in paras]
    assert "참석: 김민수, 이영희" in texts and "주요 안건: 예산 검토" in texts and "접근 방식 정리: 분기별로 본다" in texts and "다음에 볼 날: 다음 주 월요일" in texts
    assert "☐ 견적 받기 — 재무팀 · 김민수 · 기한 금요일" in texts and "대화 내용" not in texts


def test_minutes_docx_with_empty_label_does_not_crash(full, tmp_path):
    f = full / "summaries" / "minutes.json"
    data = json.loads(f.read_text(encoding="utf-8"))
    data["minutes"]["topics"][0]["points"].append({"label": "", "details": [_line("이름표 없는 줄", [])]})
    f.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    paras = _docx_paras(export("docx", tmp_path / "o.docx", "--summary", "minutes", "--parts", ""))
    assert ("List Bullet 2", "이름표 없는 줄") in paras


def test_export_parts_only_and_no_summary(full, tmp_path):
    md = export("md", tmp_path / "o.md", "--parts", "memo").read_text(encoding="utf-8")
    assert md.endswith("# 주간회의\n\n---\n\n## 메모\n\n첫 메모\n\n둘째\n") and "안녕하십니까" not in md and "요약" not in md
    md = export("md", tmp_path / "p.md", "--summary", "brief", "--parts", "").read_text(encoding="utf-8")
    assert "포인트 하나" in md and "안녕하십니까" not in md and "## 메모" not in md


def test_export_select_memo_and_recorded_notes(meeting, tmp_path):
    memo = {"schema_version": 2, "kind": "memo", "meeting_id": MID, "title": "메모", "created_at": "2026-10-02T14:00:00+09:00", "body": "가\n\n나",
            "audio": {"file": None}, "speakers": {}, "segments": [{"id": "s0001", "start": 0, "end": 0, "line": 1, "text": "가"}, {"id": "s0002", "start": 0, "end": 0, "line": 3, "text": "나"}]}
    save_transcript(meeting, memo)
    (meeting / "memos.json").unlink()
    md = export("md", tmp_path / "m.md", "--parts", "transcript,memo,marks").read_text(encoding="utf-8")
    assert md.endswith("# 메모\n\n---\n\n## 메모 내용\n\n가\n\n나\n") and "kind: memo" in md
    assert _docx_paras(export("docx", tmp_path / "m.docx", "--parts", "transcript"))[-3:] == [("Heading 1", "메모 내용"), ("Normal", "가"), ("Normal", "나")]

    recorded = {**_transcript(), "state": "recorded", "segments": []}
    del recorded["asr"]
    save_transcript(meeting, recorded)
    (meeting / "memos.json").write_text('[{"t": 5.0, "kind": "memo", "text": "중요"}]', encoding="utf-8")
    md = export("md", tmp_path / "r.md", "--parts", "transcript,marks").read_text(encoding="utf-8")
    assert "아직 텍스트로 변환하지 않았습니다." in md and "## 대화 내용" not in md and md.endswith("## 녹음 중 메모\n\n[00:00:05] (메모) 중요\n")
    texts = [x for _, x in _docx_paras(export("docx", tmp_path / "r.docx", "--parts", "transcript,marks"))]
    assert "대화 내용" not in texts and texts[-2:] == ["녹음 중 메모", "00:00:05  (메모) 중요"]


def test_export_select_errors(full, tmp_path):
    from tractus import summaries as sums
    from tractus.export import NO_SELECT

    for fmt in ("txt", "srt", "vtt", "csv", "bundle"):  # 다른 형식에 --parts를 주면 거절한다 (--summary 하나는 예전처럼 무시)
        r = invoke("export", MID, "--format", fmt, "--out", tmp_path / f"o.{fmt}", "--parts", "transcript")
        assert r.exit_code == 2 and NO_SELECT in r.output and not (tmp_path / f"o.{fmt}").exists()
    assert invoke("export", MID, "--format", "md", "--out", tmp_path / "a.md", "--parts", "transcript,audio").exit_code == 2
    r = invoke("export", MID, "--format", "md", "--out", tmp_path / "b.md", "--summary", "brief", "--summary", "todo", "--parts", "transcript")
    assert r.exit_code == 2 and sums.NO_SUCH_SUMMARY in r.output
    assert invoke("export", MID, "--format", "md", "--out", tmp_path / "c.md", "--summary", "brief", "--summary", "meeting").exit_code == 2  # 여럿은 --parts와 함께
    assert not any((tmp_path / f).exists() for f in ("a.md", "b.md", "c.md"))
